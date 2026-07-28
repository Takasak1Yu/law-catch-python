import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from db_manager import APP_DATA_DIR, ensure_app_data_dir


def generate_results_excel(results: list[dict]) -> str | None:
    """根据爬取结果生成Excel表格，异常行标红，返回文件路径"""
    ensure_app_data_dir()
    excel_dir = os.path.join(APP_DATA_DIR, "excel_reports")
    os.makedirs(excel_dir, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "爬取结果"

    headers = ["序号", "网站名称", "状态", "消息", "新通知数", "新通知标题", "链接", "摘要"]
    ws.append(headers)

    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=11, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    error_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    error_font = Font(color="9C0006")
    normal_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    normal_font = Font(color="006100")

    row_idx = 2
    status_map = {
        "error": ("异常", error_fill, error_font),
        "not_initialized": ("未初始化", error_fill, error_font),
        "normal": ("正常", normal_fill, normal_font),
        "update": ("有更新", None, None),
    }

    for result in results:
        site_name = result["site_name"]
        status = result["status"]
        message = result.get("message", "")
        new_items = result.get("new_items", [])
        summaries = result.get("summaries", {})

        status_text, fill, font = status_map.get(status, (status, None, None))

        if not new_items:
            ws.append([row_idx - 1, site_name, status_text, message, 0, "", "", ""])
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if fill:
                    cell.fill = fill
                if font:
                    cell.font = font
            row_idx += 1
        else:
            for i, (title, url) in enumerate(new_items):
                summary = summaries.get(url, "")
                ws.append([
                    row_idx - 1,
                    site_name,
                    status_text,
                    message if i == 0 else "",
                    len(new_items) if i == 0 else "",
                    title,
                    url,
                    summary,
                ])
                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if fill:
                        cell.fill = fill
                    if font:
                        cell.font = font
                row_idx += 1

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 50
    ws.column_dimensions["G"].width = 50
    ws.column_dimensions["H"].width = 50

    ws.freeze_panes = "A2"

    filename = f"爬取结果_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(excel_dir, filename)
    wb.save(filepath)

    return filepath
